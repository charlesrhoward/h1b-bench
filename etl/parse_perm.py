"""Parse DOL OFLC PERM disclosure xlsx files into employer headcount filings.

PERM (Form ETA-9089) is the only DOL disclosure that reports an employer's total
headcount: EMP_NUM_PAYROLL on the 2023+ form, EMPLOYER_NUM_EMPLOYEES on the old one.
The value is self-reported per filing, so it is kept per case here and aggregated in
load_headcounts.py.

Only employer-level columns are read. Contact names, emails, and worker details in the
source files are never extracted.

Input:  data/raw/PERM_Disclosure_Data_*.xlsx
Output: data/processed/perm_FY<year>.parquet  (one row per case_number)
"""
import glob
import logging
import os
import re
import sys

import pandas as pd
from cli_log import configure_logging
from openpyxl import load_workbook

log = logging.getLogger(__name__)

RAW_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "raw")
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "processed")

# Canonical column -> possible source names (new form first, then the pre-2023 form)
COLUMN_MAP = {
    "case_number": ["CASE_NUMBER"],
    "case_status": ["CASE_STATUS"],
    "received_date": ["RECEIVED_DATE"],
    "decision_date": ["DECISION_DATE"],
    "employer_name": ["EMP_BUSINESS_NAME", "EMPLOYER_NAME"],
    "employer_fein": ["EMP_FEIN", "EMPLOYER_FEIN"],
    "employer_state": ["EMP_STATE", "EMPLOYER_STATE_PROVINCE", "EMPLOYER_STATE"],
    "naics_code": ["EMP_NAICS", "NAICS_CODE"],
    "num_employees": ["EMP_NUM_PAYROLL", "EMPLOYER_NUM_EMPLOYEES"],
    "year_commenced": ["EMP_YEAR_COMMENCED", "EMPLOYER_YEAR_COMMENCED_BUSINESS"],
}


def normalize_fein(value):
    """Digits-only 9-digit FEIN, or None when the value is missing or malformed."""
    if value is None:
        return None
    digits = re.sub(r"\D", "", str(value))
    return digits if len(digits) == 9 else None


def column_indexes(header):
    """Map each canonical column to its index in the sheet header (None if absent)."""
    upper = [str(h).strip().upper() if h is not None else "" for h in header]
    indexes = {}
    for canon, sources in COLUMN_MAP.items():
        indexes[canon] = next((upper.index(s) for s in sources if s in upper), None)
    missing = [c for c in ("case_number", "employer_name", "num_employees") if indexes[c] is None]
    if missing:
        raise ValueError(f"required PERM columns not found: {missing}")
    return indexes


def read_rows(path):
    """Stream only the mapped columns out of the first worksheet."""
    wb = load_workbook(path, read_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    indexes = column_indexes(next(rows))
    records = []
    for row in rows:
        records.append({c: (row[i] if i is not None else None) for c, i in indexes.items()})
    wb.close()
    return pd.DataFrame.from_records(records, columns=list(COLUMN_MAP))


def parse_file(path):
    fy = int(re.search(r"FY(\d{4})", os.path.basename(path)).group(1))
    df = read_rows(path)

    for c in ("case_number", "case_status", "employer_name", "employer_state"):
        df[c] = df[c].map(lambda v: re.sub(r"\s+", " ", str(v)).strip() if v is not None else None)
    df["employer_fein"] = df["employer_fein"].map(normalize_fein)
    df["received_date"] = pd.to_datetime(df["received_date"], errors="coerce").dt.date
    df["decision_date"] = pd.to_datetime(df["decision_date"], errors="coerce").dt.date
    for c in ("num_employees", "naics_code", "year_commenced"):
        df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")

    df["fiscal_year"] = fy
    return df.dropna(subset=["case_number"]).drop_duplicates(subset="case_number", keep="first")


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    only = sys.argv[1] if len(sys.argv) > 1 else None
    for f in sorted(glob.glob(os.path.join(RAW_DIR, "PERM_Disclosure_Data_*.xlsx"))):
        if only and only not in f:
            continue
        df = parse_file(f)
        tag = re.search(r"FY\d{4}(?:_Q\d)?", os.path.basename(f)).group(0)
        out_path = os.path.join(OUT_DIR, f"perm_{tag}.parquet")
        df.to_parquet(out_path, index=False)
        with_count = df["num_employees"].notna().sum()
        log.info(f"{os.path.basename(f)}: {len(df):,} cases, {with_count:,} with headcount -> {out_path}")


if __name__ == "__main__":
    configure_logging()
    main()
